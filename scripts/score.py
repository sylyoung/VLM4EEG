#!/usr/bin/env python3
"""Recompute the four CHSZ metrics from the two released per-trial workbooks."""
import argparse
import json
from pathlib import Path

import openpyxl

from raicl import binary_metrics, parse_decision


def load_workbook_rows(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError("A released workbook must contain exactly one model sheet")
    sheet = workbook.active
    cells = sheet.iter_rows(values_only=True)
    header = next(cells)
    required = ["patient", "trial", "true_label", "predicted_label", "reference_image_1_seizure", "reference_image_2_seizure", "reference_image_3_non_seizure", "reference_image_4_non_seizure", "response"]
    if list(header) != required:
        raise ValueError("Unexpected workbook schema")
    rows, seen = [], set()
    for values in cells:
        row = dict(zip(header, values))
        key = row["patient"], row["trial"]
        if key in seen:
            raise ValueError("Repeated trial in workbook")
        seen.add(key)
        if row["true_label"] not in ("SEIZURE", "NON-SEIZURE") or row["predicted_label"] not in ("SEIZURE", "NON-SEIZURE", "UNANSWERED"):
            raise ValueError("Invalid label")
        prediction = {"SEIZURE": 1, "NON-SEIZURE": 0, "UNANSWERED": -1}[row["predicted_label"]]
        if prediction != parse_decision(row["response"]):
            raise ValueError("Stored label differs from the response's final decision")
        label = int(row["true_label"] == "SEIZURE")
        if not label and row["trial"] % 10:
            raise ValueError("Non-seizure query lies outside the CHSZ evaluation list")
        rows.append({"subject_id": row["patient"], "true_label": label, "pred_label": prediction, "error": ""})
    workbook.close()
    if len(rows) != 2783 or len({r["subject_id"] for r in rows}) != 27:
        raise ValueError("The CHSZ release requires 2,783 trials from 27 patients")
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbooks", nargs="+", type=Path)
    args = parser.parse_args()
    for path in args.workbooks:
        average, patients = binary_metrics(load_workbook_rows(path))
        print(json.dumps({"file": path.name, "patients": len(patients), "trials": sum(r["trials"] for r in patients), "metrics_percent": {k: round(v, 2) for k,v in average.items()}}, indent=2))


if __name__ == "__main__":
    main()
