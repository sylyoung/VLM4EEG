"""Offline regression checks for installation, documentation and the released artifacts."""
import hashlib
from html.parser import HTMLParser
from importlib import metadata
import json
import os
from pathlib import Path
import re
import subprocess
import sys
from urllib.parse import unquote, urlsplit

import pytest

from raicl import binary_metrics, prompt_from_paths
from raicl.paths import DEFAULT_PROMPT, PROMPTS, TRAINING_MANIFEST
from raicl.workflows.gemini import fingerprint
from raicl.workflows.score import load_workbook_rows

ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = {
    "CHSZ_Gemini-3-Flash_RAICL.xlsx": [82.28, 78.09, 86.48, 74.25],
    "CHSZ_Qwen3.5-9B_RAICL-finetuned.xlsx": [82.84, 74.41, 91.27, 78.00],
}
WORKFLOWS = ["render", "embed", "prepare", "gemini", "train", "qwen", "score"]


@pytest.mark.parametrize("name,expected", WORKBOOKS.items())
def test_released_metrics(name, expected):
    metrics, patients = binary_metrics(load_workbook_rows(ROOT / "results" / name))
    assert len(patients) == 27
    assert sum(row["trials"] for row in patients) == 2783
    assert [round(value, 2) for value in metrics.values()] == expected


def test_benchmark_artifact_checksums():
    checksums = json.loads((ROOT / "tests/fixtures/release_checksums.json").read_text())
    for name, expected in checksums.items():
        assert hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == expected, name


def test_both_workbooks_have_the_same_trial_list():
    import openpyxl
    identities = []
    for name in WORKBOOKS:
        workbook = openpyxl.load_workbook(ROOT / "results" / name, read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        identities.append({(row[0], row[1], row[2]) for row in rows})
        workbook.close()
    assert len(identities[0]) == 2783
    assert identities[0] == identities[1]


def test_packaged_resources_do_not_depend_on_working_directory(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    assert DEFAULT_PROMPT.is_file() and TRAINING_MANIFEST.is_file()
    assert len(json.loads(TRAINING_MANIFEST.read_text())) == 27
    query = "data/VLM/sub_00/class_1/trial_30_VLM.png"
    refs = ["data/VLM/sub_00/class_0/trial_10_VLM.png",
            "data/VLM/sub_00/class_0/trial_20_VLM.png",
            "data/VLM/sub_01/class_1/trial_40_VLM.png",
            "data/VLM/sub_02/class_1/trial_50_VLM.png"]
    _, images = prompt_from_paths(query, refs)
    assert images == [refs[3], refs[2], refs[0], refs[1], query]


@pytest.mark.parametrize("workflow", WORKFLOWS)
def test_module_help_is_offline(workflow, tmp_path):
    result = subprocess.run(
        [sys.executable, "-m", "raicl.workflows." + workflow, "--help"],
        cwd=tmp_path, capture_output=True, text=True, timeout=45,
    )
    assert result.returncode == 0, result.stderr
    assert "usage:" in result.stdout


def test_installed_console_entry_points():
    entries = {entry.name: entry for entry in metadata.distribution("raicl-eeg").entry_points if entry.group == "console_scripts"}
    expected = {"raicl-" + name for name in WORKFLOWS}
    assert expected == set(entries)
    for entry in entries.values():
        assert callable(entry.load())


def test_core_import_does_not_load_model_runtimes(tmp_path):
    code = "import sys; import raicl; assert not {'torch', 'transformers', 'llamafactory', 'google.genai', 'clip'} & sys.modules.keys()"
    result = subprocess.run([sys.executable, "-c", code], cwd=tmp_path, capture_output=True, text=True, timeout=45)
    assert result.returncode == 0, result.stderr


def test_fingerprint_is_sha256_on_supported_python_versions(tmp_path):
    path = tmp_path / "fingerprint.txt"
    path.write_bytes(b"benchmark prompt\n" * 100000)
    assert fingerprint(path) == hashlib.sha256(path.read_bytes()).hexdigest()


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.targets = []

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        for name in ("href", "src"):
            if name in values:
                self.targets.append(values[name])
        if tag == "img":
            assert values.get("alt"), "README images need descriptive alternative text"


def heading_ids(text):
    result = set()
    for heading in re.findall(r"^#{1,6}\s+(.+)$", text, re.M):
        slug = re.sub(r"[^\w\- ]", "", heading.strip().lower()).replace(" ", "-")
        result.add(slug)
    return result


def test_documentation_links():
    documents = [ROOT / "README.md", *ROOT.glob("docs/*.md"), *ROOT.glob("examples/*.md"), *ROOT.glob("results/*.md"), *ROOT.glob("prompts/*.md")]
    for document in documents:
        text = document.read_text()
        parser = Links()
        parser.feed(text)
        targets = re.findall(r"\]\(([^\s)]+)\)", text) + parser.targets
        for target in targets:
            parts = urlsplit(target)
            if parts.scheme or parts.netloc:
                continue
            path = (document.parent / unquote(parts.path)).resolve() if parts.path else document
            assert path.exists(), f"{document.name}: missing {target}"
            if parts.fragment and path.suffix == ".md":
                assert unquote(parts.fragment) in heading_ids(path.read_text()), f"{document.name}: missing anchor {target}"


def test_private_and_generated_files_are_ignored():
    names = ["data/CHSZ/X.npy", "models/adapter/model.safetensors", "outputs/response.csv",
             ".env", ".env.local", "secret.pem", "submission/manuscript.tex",
             "review/comments.txt", "reports/local.json", "results/unreleased.xlsx"]
    result = subprocess.run(["git", "check-ignore", "--stdin"], input="\n".join(names), cwd=ROOT, text=True, capture_output=True, check=True)
    assert set(result.stdout.splitlines()) == set(names)
    for name in WORKBOOKS:
        assert subprocess.run(["git", "check-ignore", "--no-index", "results/" + name], cwd=ROOT, capture_output=True).returncode == 1


def test_prompt_directory_names_distinguish_experimental_content():
    assert {p.name for p in (PROMPTS / "binary").glob("*.txt")} == {"raicl.txt", "base.txt", "reasoning.txt", "random_examples.txt"}
    assert {p.name for p in (PROMPTS / "experimental").glob("*.txt")} == {"decision_only.txt", "concrete_wording.txt"}
